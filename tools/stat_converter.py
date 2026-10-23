"""Stat Template converter (v0.2; v0.3: Stat Template v1.3 headers + Economy tab).

Reads docs/Stat_Template_v1.3.xlsx and writes the Luau config modules:
    src/shared/Config/Schema.luau        tab/field schema + enum values
    src/shared/Config/Data/<Tab>.luau    one record per filled row

Usage (from the repo root):
    python tools/stat_converter.py [path/to/template.xlsx]

Rules (deterministic, every branch explicit):
- Row 1 of every tab must match the expected headers exactly, in order. Any
  difference stops the run, so a template change always needs a matching edit here.
- A row whose ID cell is blank is skipped (unfilled template row).
- A row whose ID starts with "REF-" is skipped (Clash of Clans reference row).
- A blank cell becomes nil (Legend: blank = not yet decided).
- Enum values come from the template's own dropdown lists. A value outside the
  list stops the run with the cell coordinate.
- Y/N becomes true/false.
- Every record carries _Row (its sheet row) so validation errors trace to a cell.

Generated files are overwritten on every run. Do not edit them by hand.
"""
import re
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = REPO / "docs" / "Stat_Template_v1.3.xlsx"
OUT_DIR = REPO / "src" / "shared" / "Config"

# Field types:
#   id            primary key of a base tab (pattern checked at server start)
#   key           primary key of a settings tab (identifier, e.g. StartingCoins)
#   ref:<Tab>     ID that must exist in <Tab>
#   string        free text
#   int           whole number
#   number        any number
#   bool          Y/N
#   enum:<Name>   value from the dropdown list <Name>
#   footprint     "<W>x<D>" in tiles
#   idlist:<Tab>  comma-separated IDs that must exist in <Tab>
TABS = [
    {
        "Sheet": "Units", "Module": "Units", "Kind": "base", "IdPrefix": "U",
        "Fields": [
            ("Unit ID", "UnitId", "id"),
            ("Name", "Name", "string"),
            ("Faction", "Faction", "enum:Faction"),
            ("Class (Troop/Hero)", "Class", "enum:Class"),
            ("Town Hall unlock", "TownHallUnlock", "int"),
            ("Housing space", "HousingSpace", "int"),
            ("Movement domain", "MovementDomain", "enum:MovementDomain"),
            ("Can target", "CanTarget", "enum:TargetDomain"),
            ("Preferred target", "PreferredTarget", "enum:PreferredTarget"),
            ("Damage multiplier vs preferred target", "PreferredTargetDamageMultiplier", "number"),
            ("Attack type", "AttackType", "enum:AttackType"),
            ("Splash radius (tiles)", "SplashRadius", "number"),
            ("Range (tiles)", "Range", "number"),
            ("Attack speed (s)", "AttackSpeed", "number"),
            ("Movement speed", "MovementSpeed", "number"),
            ("Wall interaction", "WallInteraction", "enum:WallInteraction"),
            ("Special behavior", "SpecialBehavior", "string"),
            ("Notes", "Notes", "string"),
        ],
    },
    {
        "Sheet": "Unit Levels", "Module": "UnitLevels", "Kind": "levels", "Parent": "Units",
        "Fields": [
            ("Unit ID", "UnitId", "ref:Units"),
            ("Level", "Level", "int"),
            ("Hitpoints", "Hitpoints", "number"),
            ("Damage per second", "DamagePerSecond", "number"),
            ("Upgrade cost (Coins)", "UpgradeCost", "number"),
            ("Upgrade time (s)", "UpgradeTimeSeconds", "number"),
            ("Required Town Hall", "RequiredTownHall", "int"),
            ("Required research building level", "RequiredLaboratoryLevel", "int"),
            ("Value status", "ValueStatus", "enum:ValueStatus"),
            ("Notes", "Notes", "string"),
        ],
    },
    {
        "Sheet": "Heroes", "Module": "Heroes", "Kind": "extension", "Parent": "Units",
        "Fields": [
            ("Unit ID (from Units)", "UnitId", "ref:Units"),
            ("Element", "Element", "enum:Element"),
            ("Hero Hall unlock level", "HeroHallUnlock", "int"),
            ("Ability name", "AbilityName", "string"),
            ("Ability effect", "AbilityEffect", "string"),
            ("Ability trigger", "AbilityTrigger", "string"),
            ("Equipment slots", "EquipmentSlots", "int"),
            ("Defends at home base (Y/N)", "DefendsAtHome", "bool"),
            ("Awakening max tier", "AwakeningMaxTier", "int"),
            ("Notes", "Notes", "string"),
        ],
    },
    {
        "Sheet": "Defenses", "Module": "Defenses", "Kind": "base", "IdPrefix": "D",
        "Fields": [
            ("Defense ID", "DefenseId", "id"),
            ("Name", "Name", "string"),
            ("Faction", "Faction", "enum:Faction"),
            ("Town Hall unlock", "TownHallUnlock", "int"),
            ("Footprint (tiles, e.g. 3x3)", "Footprint", "footprint"),
            ("Target domain", "TargetDomain", "enum:TargetDomain"),
            ("Attack type", "AttackType", "enum:AttackType"),
            ("Splash radius (tiles)", "SplashRadius", "number"),
            ("Min range / blind spot (tiles)", "MinRange", "number"),
            ("Max range (tiles)", "MaxRange", "number"),
            ("Attack speed (s)", "AttackSpeed", "number"),
            ("Player-set modes", "PlayerSetModes", "string"),
            ("Hidden until triggered (Y/N)", "HiddenUntilTriggered", "bool"),
            ("Counts toward destruction % (Y/N)", "CountsTowardDestruction", "bool"),
            ("No-deploy zone (tiles beyond footprint)", "NoDeployZone", "number"),
            ("Notes", "Notes", "string"),
        ],
    },
    {
        "Sheet": "Defense Levels", "Module": "DefenseLevels", "Kind": "levels", "Parent": "Defenses",
        "Fields": [
            ("Defense ID", "DefenseId", "ref:Defenses"),
            ("Level", "Level", "int"),
            ("Hitpoints", "Hitpoints", "number"),
            ("Damage per second", "DamagePerSecond", "number"),
            ("Upgrade cost (Coins)", "UpgradeCost", "number"),
            ("Upgrade time (s)", "UpgradeTimeSeconds", "number"),
            ("Required Town Hall", "RequiredTownHall", "int"),
            ("Count allowed at this Town Hall", "CountAllowed", "int"),
            ("Value status", "ValueStatus", "enum:ValueStatus"),
            ("Notes", "Notes", "string"),
        ],
    },
    {
        "Sheet": "Traps", "Module": "Traps", "Kind": "base", "IdPrefix": "T",
        "Fields": [
            ("Trap ID", "TrapId", "id"),
            ("Name", "Name", "string"),
            ("Faction", "Faction", "enum:Faction"),
            ("Town Hall unlock", "TownHallUnlock", "int"),
            ("Trigger domain", "TriggerDomain", "enum:TargetDomain"),
            ("Trigger radius (tiles)", "TriggerRadius", "number"),
            ("Target rule", "TargetRule", "string"),
            ("Effect", "Effect", "string"),
            ("Housing capacity", "HousingCapacity", "int"),
            ("Notes", "Notes", "string"),
        ],
    },
    {
        "Sheet": "Trap Levels", "Module": "TrapLevels", "Kind": "levels", "Parent": "Traps",
        "Fields": [
            ("Trap ID", "TrapId", "ref:Traps"),
            ("Level", "Level", "int"),
            ("Damage", "Damage", "number"),
            ("Upgrade cost (Coins)", "UpgradeCost", "number"),
            ("Upgrade time (s)", "UpgradeTimeSeconds", "number"),
            ("Required Town Hall", "RequiredTownHall", "int"),
            ("Value status", "ValueStatus", "enum:ValueStatus"),
            ("Notes", "Notes", "string"),
        ],
    },
    {
        "Sheet": "Buildings", "Module": "Buildings", "Kind": "base", "IdPrefix": "B",
        "Fields": [
            ("Building ID", "BuildingId", "id"),
            ("Name", "Name", "string"),
            ("Faction", "Faction", "enum:Faction"),
            ("Building type", "BuildingType", "enum:BuildingType"),
            ("Resource type", "ResourceType", "enum:ResourceType"),
            ("Town Hall unlock", "TownHallUnlock", "int"),
            ("Footprint (tiles, e.g. 3x3)", "Footprint", "footprint"),
            ("Counts toward destruction % (Y/N)", "CountsTowardDestruction", "bool"),
            ("No-deploy zone (tiles beyond footprint)", "NoDeployZone", "number"),
            ("Lootable (Y/N)", "Lootable", "bool"),
            ("Notes", "Notes", "string"),
        ],
    },
    {
        "Sheet": "Building Levels", "Module": "BuildingLevels", "Kind": "levels", "Parent": "Buildings",
        "Fields": [
            ("Building ID", "BuildingId", "ref:Buildings"),
            ("Level", "Level", "int"),
            ("Hitpoints", "Hitpoints", "number"),
            ("Upgrade cost (Coins)", "UpgradeCost", "number"),
            ("Upgrade time (s)", "UpgradeTimeSeconds", "number"),
            ("Required Town Hall", "RequiredTownHall", "int"),
            ("Count allowed at this Town Hall", "CountAllowed", "int"),
            ("Resource production rate (per hour)", "ProductionPerHour", "number"),
            ("Resource capacity", "ResourceCapacity", "number"),
            ("Market sell limit (per cycle)", "MarketSellLimit", "number"),
            ("Army Camp housing capacity", "ArmyCampHousing", "int"),
            ("Max hero level", "MaxHeroLevel", "int"),
            ("Hero slots", "HeroSlots", "int"),
            ("Builders provided", "BuildersProvided", "int"),
            ("Unlocks (unit IDs, comma-separated)", "Unlocks", "idlist:Units"),
            ("Value status", "ValueStatus", "enum:ValueStatus"),
            ("Notes", "Notes", "string"),
        ],
    },
    {
        "Sheet": "Economy", "Module": "Economy", "Kind": "settings",
        "Fields": [
            ("Setting ID", "SettingId", "key"),
            ("Value", "Value", "number"),
            ("Unit", "Unit", "string"),
            ("Value status", "ValueStatus", "enum:ValueStatus"),
            ("Notes", "Notes", "string"),
        ],
    },
]

TEMPLATE_TITLE = re.compile(r"^Stat Template (v\d+\.\d+)$")
FOOTPRINT = re.compile(r"^\d+x\d+$")


class ConvertError(Exception):
    pass


def read_enums(wb):
    """Enum name -> allowed values, taken from the dropdown lists of enum columns.

    The same enum on several tabs (e.g. Faction) must have an identical list everywhere.
    """
    enums = {}
    for tab in TABS:
        ws = wb[tab["Sheet"]]
        for col, (header, key, ftype) in enumerate(tab["Fields"], start=1):
            if not ftype.startswith("enum:"):
                continue
            name = ftype.split(":", 1)[1]
            letter = openpyxl.utils.get_column_letter(col)
            found = None
            for dv in ws.data_validations.dataValidation:
                if dv.type != "list":
                    continue
                cols = {openpyxl.utils.get_column_letter(c) for rng in dv.sqref.ranges
                        for c in range(rng.min_col, rng.max_col + 1)}
                if letter in cols:
                    found = [v.strip() for v in dv.formula1.strip('"').split(",")]
                    break
            if found is None:
                raise ConvertError(f"{tab['Sheet']}!{letter}1 '{header}': no dropdown list for enum {name}")
            if name in enums and enums[name] != found:
                raise ConvertError(
                    f"{tab['Sheet']}!{letter}1 '{header}': dropdown list {found} differs from "
                    f"enum {name} {enums[name]} defined on another tab")
            enums[name] = found
    return enums


def convert_cell(value, ftype, enums, where):
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    if ftype in ("id", "key", "string") or ftype.startswith("ref:"):
        return str(value).strip()
    if ftype in ("int", "number"):
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ConvertError(f"{where}: expected a number, got {value!r}")
        if ftype == "int":
            if float(value) != int(value):
                raise ConvertError(f"{where}: expected a whole number, got {value!r}")
            return int(value)
        return int(value) if float(value) == int(value) else float(value)
    if ftype == "bool":
        text = str(value).strip()
        if text == "Y":
            return True
        if text == "N":
            return False
        raise ConvertError(f"{where}: expected Y or N, got {value!r}")
    if ftype.startswith("enum:"):
        name = ftype.split(":", 1)[1]
        text = str(value).strip()
        if text not in enums[name]:
            raise ConvertError(f"{where}: {text!r} is not in {name} {enums[name]}")
        return text
    if ftype == "footprint":
        text = str(value).strip()
        if not FOOTPRINT.match(text):
            raise ConvertError(f"{where}: expected <W>x<D> (e.g. 3x3), got {value!r}")
        return text
    if ftype.startswith("idlist:"):
        return [part.strip() for part in str(value).split(",") if part.strip() != ""]
    raise ConvertError(f"{where}: unknown field type {ftype!r}")


def read_tab(wb, tab, enums):
    ws = wb[tab["Sheet"]]
    expected = [header for header, _, _ in tab["Fields"]]
    actual = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    while actual and actual[-1] is None:
        actual.pop()
    if actual != expected:
        raise ConvertError(f"{tab['Sheet']}: header row does not match.\n  expected {expected}\n  found    {actual}")

    records, skipped_blank, skipped_ref = [], 0, 0
    for row in range(2, ws.max_row + 1):
        first = ws.cell(row, 1).value
        if first is None or str(first).strip() == "":
            skipped_blank += 1
            continue
        if str(first).strip().startswith("REF-"):
            skipped_ref += 1
            continue
        record = {"_Row": row}
        for col, (header, key, ftype) in enumerate(tab["Fields"], start=1):
            where = f"{tab['Sheet']}!{ws.cell(row, col).coordinate} '{header}'"
            record[key] = convert_cell(ws.cell(row, col).value, ftype, enums, where)
        records.append(record)
    return records, skipped_blank, skipped_ref


def lua_value(value):
    if value is None:
        return "nil"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, list):
        return "{ " + ", ".join(lua_value(v) for v in value) + " }" if value else "{}"
    text = str(value).replace("\\", "\\\\").replace('"', '\\"').replace("\r", "").replace("\n", "\\n")
    return f'"{text}"'


def header_comment(template_name, version):
    return (f"-- GENERATED by tools/stat_converter.py from {template_name} (Stat Template {version}).\n"
            f"-- Do not edit by hand: edit the template and re-run the converter.\n")


def write_data(tab, records, template_label, version):
    lines = [header_comment(template_label, version), "return {"]
    for record in records:
        lines.append("\t{")
        lines.append(f"\t\t_Row = {record['_Row']},")
        for _, key, _ in tab["Fields"]:
            if record[key] is not None:
                lines.append(f"\t\t{key} = {lua_value(record[key])},")
        lines.append("\t},")
    lines.append("}")
    return "\n".join(lines) + "\n"


def write_schema(enums, version, template_label):
    lines = [header_comment(template_label, version), "return {", f'\tTemplateVersion = "{version}",', "\tEnums = {"]
    for name in sorted(enums):
        lines.append(f"\t\t{name} = {{ " + ", ".join(lua_value(v) for v in enums[name]) + " },")
    lines.append("\t},")
    lines.append("\tTabs = {")
    for tab in TABS:
        lines.append("\t\t{")
        lines.append(f'\t\t\tSheet = "{tab["Sheet"]}",')
        lines.append(f'\t\t\tModule = "{tab["Module"]}",')
        lines.append(f'\t\t\tKind = "{tab["Kind"]}",')
        if "IdPrefix" in tab:
            lines.append(f'\t\t\tIdPrefix = "{tab["IdPrefix"]}",')
        if "Parent" in tab:
            lines.append(f'\t\t\tParent = "{tab["Parent"]}",')
        lines.append("\t\t\tFields = {")
        for header, key, ftype in tab["Fields"]:
            lines.append(f"\t\t\t\t{{ Key = {lua_value(key)}, Header = {lua_value(header)}, Type = {lua_value(ftype)} }},")
        lines.append("\t\t\t},")
        lines.append("\t\t},")
    lines.append("\t},")
    lines.append("}")
    return "\n".join(lines) + "\n"


def main(argv):
    template = (Path(argv[1]) if len(argv) > 1 else DEFAULT_TEMPLATE).resolve()
    wb = openpyxl.load_workbook(template)
    if template.is_relative_to(REPO):
        template_label = template.relative_to(REPO).as_posix()
    else:
        template_label = template.as_posix()

    title = wb["Legend"]["A1"].value
    match = TEMPLATE_TITLE.match(str(title))
    if not match:
        raise ConvertError(f"Legend!A1: expected 'Stat Template vX.Y', got {title!r}")
    version = match.group(1)

    enums = read_enums(wb)
    outputs = {OUT_DIR / "Schema.luau": write_schema(enums, version, template_label)}
    summary = []
    for tab in TABS:
        records, skipped_blank, skipped_ref = read_tab(wb, tab, enums)
        outputs[OUT_DIR / "Data" / f"{tab['Module']}.luau"] = write_data(tab, records, template_label, version)
        summary.append(f"  {tab['Sheet']:<16} {len(records):>4} records  "
                       f"({skipped_blank} blank-ID rows, {skipped_ref} reference rows skipped)")

    # Write only after every tab converted cleanly, so a failed run leaves no partial output.
    (OUT_DIR / "Data").mkdir(parents=True, exist_ok=True)
    for path, text in outputs.items():
        path.write_text(text, encoding="utf-8", newline="\n")

    print(f"Stat Template {version} ({template_label}) -> {OUT_DIR.relative_to(REPO).as_posix()}")
    print("\n".join(summary))


if __name__ == "__main__":
    try:
        main(sys.argv)
    except ConvertError as err:
        print(f"ERROR: {err}", file=sys.stderr)
        sys.exit(1)
